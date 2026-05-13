"""
Build the Bennett Cookbook recipe database exports.

Outputs:
  exports/the_bennett_cookbook_recipe_database.xlsx  (multi-tab)
  exports/the_bennett_cookbook_recipe_database.csv
  exports/recipes_export.json

Run from repo root:
  python -m src.build_database
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent.parent
sys.path.insert(0, str(BASE))

from src.models import FullMeal, Recipe


# ─── Load / validate ────────────────────────────────────────────────────────

def load_recipes(path: Path) -> list[Recipe]:
    raw = json.loads(path.read_text(encoding="utf-8"))
    return [Recipe.model_validate(r) for r in raw]


def load_saved_external(path: Path) -> list[Recipe]:
    if not path.exists():
        return []
    raw = json.loads(path.read_text(encoding="utf-8"))
    return [Recipe.model_validate(r) for r in raw]


def load_full_meals(path: Path) -> list[FullMeal]:
    raw = json.loads(path.read_text(encoding="utf-8"))
    return [FullMeal.model_validate(m) for m in raw]


# ─── Flatten ────────────────────────────────────────────────────────────────

FLAT_COLUMN_ORDER = [
    "recipe_id", "recipe_name", "record_type", "category", "cuisine",
    "difficulty", "serves", "prep_time_minutes", "cook_time_minutes",
    "total_time_minutes", "kid_friendly", "nathan_rating", "leah_rating",
    "nina_rating", "weeknight_score", "entertaining_score",
    "family_dinner_candidate", "entertaining_candidate", "try_soon",
    "tags", "pairings", "nathan_tweaks", "substitutions",
    "full_meal_id", "source_name", "source_url", "source_author",
    "date_saved", "copyright_status", "adaptation_status",
    "planned_ingredient_changes", "planned_method_changes",
    "image_file", "source_notes",
    "ingredients", "method",
]


def recipe_to_flat_dict(r: Recipe) -> dict:
    d = r.model_dump()
    d["difficulty"] = r.difficulty.value
    d["record_type"] = r.record_type.value
    d["adaptation_status"] = r.adaptation_status.value if r.adaptation_status else ""
    d["ingredients"] = " | ".join(r.ingredients)
    d["method"] = " | ".join(r.method)
    d["pairings"] = ", ".join(r.pairings)
    d["tags"] = ", ".join(r.tags)
    return d


def build_df(recipes: list[Recipe]) -> pd.DataFrame:
    rows = [recipe_to_flat_dict(r) for r in recipes]
    df = pd.DataFrame(rows)
    cols = [c for c in FLAT_COLUMN_ORDER if c in df.columns]
    extra = [c for c in df.columns if c not in cols]
    return df[cols + extra]


# ─── XLSX styling ───────────────────────────────────────────────────────────

HEADER_FILL   = PatternFill("solid", fgColor="C0392B")
HEADER_FONT   = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
ALT_FILL_1    = PatternFill("solid", fgColor="FAF8F5")
ALT_FILL_2    = PatternFill("solid", fgColor="F0EBE3")
HIDDEN_COLS   = {"ingredients", "method"}


def style_worksheet(ws: object) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        fill = ALT_FILL_1 if row_idx % 2 == 0 else ALT_FILL_2
        for cell in row:
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for col_idx, col in enumerate(ws.columns, start=1):
        header_val = ws.cell(row=1, column=col_idx).value or ""
        max_len = max(
            (len(str(c.value or "")) for c in col if c.row > 1),
            default=len(str(header_val)),
        )
        width = min(max(max_len + 2, len(str(header_val)) + 4), 50)
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width
        if str(header_val).lower() in HIDDEN_COLS:
            ws.column_dimensions[col_letter].hidden = True

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


# ─── XLSX export ────────────────────────────────────────────────────────────

def write_xlsx(
    personal_recipes: list[Recipe],
    saved_external: list[Recipe],
    full_meals: list[FullMeal],
    out_path: Path,
) -> None:
    all_recipes = personal_recipes + saved_external
    all_df = build_df(all_recipes)

    def filt(record_type: str) -> pd.DataFrame:
        return all_df[all_df["record_type"] == record_type].copy()

    out_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:

        def write_tab(df: pd.DataFrame, sheet_name: str) -> None:
            if df.empty:
                df = pd.DataFrame(columns=all_df.columns)
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            style_worksheet(writer.book[sheet_name])

        write_tab(all_df.sort_values("recipe_name"), "All Recipes")
        write_tab(filt("personal_recipe").sort_values("recipe_name"), "Personal Recipes")
        write_tab(filt("external_saved_reference").sort_values("date_saved", ascending=False), "External References")
        write_tab(filt("adapted_bennett_recipe").sort_values("recipe_name"), "Adapted Recipes")

        # Weeknight Wins
        wn = all_df[pd.to_numeric(all_df["weeknight_score"], errors="coerce") >= 7].copy()
        wn = wn.sort_values("weeknight_score", ascending=False)
        write_tab(wn, "Weeknight Wins")

        # Entertaining
        ent = all_df[pd.to_numeric(all_df["entertaining_score"], errors="coerce") >= 7].copy()
        ent = ent.sort_values("entertaining_score", ascending=False)
        write_tab(ent, "Entertaining")

        # Full Meals
        meal_rows = []
        for m in full_meals:
            meal_rows.append({
                "full_meal_id": m.full_meal_id,
                "meal_name": m.meal_name,
                "description": m.description or "",
                "recipe_ids": " | ".join(m.recipe_ids),
                "serves": m.serves,
                "total_time_estimate": m.total_time_estimate or "",
                "occasions": ", ".join(m.occasions),
                "notes": m.notes or "",
            })
        meals_df = pd.DataFrame(meal_rows).sort_values("meal_name")
        write_tab(meals_df, "Full Meals")

        # Tag Index
        tag_rows = []
        for r in all_recipes:
            for tag in r.tags:
                tag_rows.append({
                    "tag": tag,
                    "recipe_id": r.recipe_id,
                    "recipe_name": r.recipe_name,
                    "category": r.category,
                    "difficulty": r.difficulty.value,
                })
        tag_df = pd.DataFrame(tag_rows).sort_values(["tag", "recipe_name"])
        write_tab(tag_df, "Tag Index")

    print(f"  XLSX → {out_path}")


# ─── CSV export ─────────────────────────────────────────────────────────────

def write_csv(recipes: list[Recipe], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    df = build_df(recipes)
    df.to_csv(out_path, index=False, encoding="utf-8-sig")
    print(f"  CSV  → {out_path}")


# ─── JSON export ────────────────────────────────────────────────────────────

def write_json(recipes: list[Recipe], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    data = [r.model_dump() for r in recipes]
    for d in data:
        d["difficulty"] = d["difficulty"].value if hasattr(d["difficulty"], "value") else d["difficulty"]
        d["record_type"] = d["record_type"].value if hasattr(d["record_type"], "value") else d["record_type"]
        if d.get("save_type"):
            d["save_type"] = d["save_type"].value if hasattr(d["save_type"], "value") else d["save_type"]
        if d.get("adaptation_status"):
            d["adaptation_status"] = d["adaptation_status"].value if hasattr(d["adaptation_status"], "value") else d["adaptation_status"]
    out_path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"  JSON → {out_path}")


# ─── Entry point ────────────────────────────────────────────────────────────

def main() -> None:
    print("Bennett Cookbook — Building database exports...")

    personal = load_recipes(BASE / "data" / "recipes.json")
    external = load_saved_external(BASE / "data" / "saved_external.json")
    meals    = load_full_meals(BASE / "data" / "full_meals.json")

    all_recipes = personal + external

    write_xlsx(personal, external, meals, BASE / "exports" / "the_bennett_cookbook_recipe_database.xlsx")
    write_csv(all_recipes, BASE / "exports" / "the_bennett_cookbook_recipe_database.csv")
    write_json(all_recipes, BASE / "exports" / "recipes_export.json")

    print(f"\nDone. {len(personal)} personal recipes, {len(external)} external references.")


if __name__ == "__main__":
    main()
